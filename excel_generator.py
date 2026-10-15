import io
import os
import re
import zipfile
import unicodedata
from copy import copy

import openpyxl
import pandas as pd

from config import (
    TEMPLATE_PATH,
    CONFIG_INDICADORES,
    MAPEO_COLUMNAS,
    CELDA_TITULO,
    MESES_COLUMNAS,
    HOJA_GDR,
    HOJA_UP_INDICADOR_1,
    FILA_INICIO_UP,
)
from data_processor import procesar_datos_por_cite


def normalizar_nombre_archivo(nombre):
    texto = unicodedata.normalize("NFKD", str(nombre)).encode("ascii", "ignore").decode("ascii")
    texto = texto.replace(" ", "_")
    texto = re.sub(r"[^A-Za-z0-9_-]", "_", texto)
    return re.sub(r"_+", "_", texto).strip("_")


def _valor_excel(valor):
    if valor is None:
        return None
    if isinstance(valor, str):
        return valor
    return valor


def _obtener_meta(df_metas, cite, columna):
    coincidencia = df_metas[df_metas["CITE"].astype(str).str.strip() == str(cite).strip()]
    if len(coincidencia) != 1:
        raise ValueError(
            f"Se esperaba una fila de metas para '{cite}' y se encontraron {len(coincidencia)}."
        )
    valor = pd.to_numeric(coincidencia.iloc[0][columna], errors="coerce")
    if pd.isna(valor):
        raise ValueError(f"La meta '{columna}' de '{cite}' no es numérica.")
    return float(valor)


def _copiar_estilo_fila(ws, fila_origen, fila_destino, columnas=4):
    for col in range(1, columnas + 1):
        src = ws.cell(row=fila_origen, column=col)
        dst = ws.cell(row=fila_destino, column=col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
    if ws.row_dimensions[fila_origen].height is not None:
        ws.row_dimensions[fila_destino].height = ws.row_dimensions[fila_origen].height


def _llenar_listado_up(ws_up, df_up, cite):
    detalle = df_up[df_up["CITE"].astype(str).str.strip() == str(cite).strip()].copy()
    if detalle.empty:
        raise ValueError(f"No existen registros en el listado UP para '{cite}'.")

    columnas = [
        "N°",
        "RUC/DNI/CE/OTROS",
        "RAZON SOCIAL / APELLIDOS Y NOMBRES",
        "CATEGORIA",
    ]

    # Las filas 4 y 5 de la plantilla son filas modelo con estilos alternados.
    for i, (_, registro) in enumerate(detalle.iterrows()):
        fila = FILA_INICIO_UP + i
        fila_modelo = 4 if i % 2 == 0 else 5
        if fila > 5:
            _copiar_estilo_fila(ws_up, fila_modelo, fila)

        for col, nombre_col in enumerate(columnas, start=1):
            valor = registro[nombre_col]
            if pd.isna(valor):
                valor = None
            elif nombre_col == "RUC/DNI/CE/OTROS":
                valor = str(valor).strip()
                if valor.endswith(".0") and valor[:-2].isdigit():
                    valor = valor[:-2]
            ws_up.cell(row=fila, column=col).value = valor

    # Limpia cualquier fila modelo o contenido residual posterior al listado.
    primera_vacia = FILA_INICIO_UP + len(detalle)
    for fila in range(primera_vacia, max(ws_up.max_row, primera_vacia) + 1):
        for col in range(1, 5):
            ws_up.cell(row=fila, column=col).value = None

    return len(detalle)


def generar_excel_bytes(df, df_metas, df_up, cite, anio, mes_inicio_ui, mes_fin_ui):
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(
            f"No se encontró la plantilla oficial: {TEMPLATE_PATH}."
        )

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    if HOJA_GDR not in wb.sheetnames or HOJA_UP_INDICADOR_1 not in wb.sheetnames:
        wb.close()
        raise ValueError(
            f"La plantilla debe contener las hojas '{HOJA_GDR}' y '{HOJA_UP_INDICADOR_1}'."
        )

    ws = wb[HOJA_GDR]
    ws_up = wb[HOJA_UP_INDICADOR_1]
    cite_mostrar = str(cite).upper()

    ws[CELDA_TITULO] = (
        f"FICHA TÉCNICA DE INDICADORES {anio} - {cite_mostrar}\n"
        f"(PERIODO {mes_inicio_ui.upper()} - {mes_fin_ui.upper()} {anio})"
    )

    df_res, _ = procesar_datos_por_cite(df, cite, anio, mes_inicio_ui, mes_fin_ui)

    for id_ind, config in CONFIG_INDICADORES.items():
        meta = _obtener_meta(df_metas, cite, config["columna_meta"])
        celda_meta = ws.cell(
            row=config["celda_meta"]["fila"],
            column=config["celda_meta"]["col"],
        )
        celda_meta.value = meta
        if id_ind == 3:
            celda_meta.number_format = "0.00%"
        elif float(meta).is_integer():
            celda_meta.value = int(meta)

        for variable, var_config in config["variables"].items():
            coincidencia = df_res[
                (df_res["Indicador"] == config["nombre"])
                & (df_res["Variable"] == variable)
            ]
            if len(coincidencia) != 1:
                wb.close()
                raise ValueError(
                    f"No se pudo identificar de forma única {config['nombre']} - {variable}."
                )

            fila_datos = coincidencia.iloc[0]
            fila_excel = var_config["fila"]
            for mes_col in MESES_COLUMNAS:
                ws.cell(
                    row=fila_excel,
                    column=MAPEO_COLUMNAS[mes_col],
                ).value = _valor_excel(fila_datos[mes_col])

        fila_resultado = config["celda_resultado"]["fila"]
        col_resultado = config["celda_resultado"]["col"]
        celda_resultado = ws.cell(row=fila_resultado, column=col_resultado)

        if config["tipo_calculo"] == "porcentaje":
            fila_fuente = df_res[
                (df_res["Indicador"] == config["nombre"])
                & (df_res["Variable"] == "Número de encuestas respondidas")
            ].iloc[0]
            resultado = fila_fuente["Valor Alcanzado"]
            if resultado == "-":
                celda_resultado.value = "-"
            elif resultado == "" or resultado is None:
                celda_resultado.value = None
            else:
                celda_resultado.value = float(resultado)
                celda_resultado.number_format = "0.00%"
        else:
            fila_fuente = df_res[df_res["Indicador"] == config["nombre"]].iloc[0]
            resultado = fila_fuente["Valor Alcanzado"]
            celda_resultado.value = None if resultado == "" else resultado

            # Solo Indicadores 1 y 2 muestran avance.
            celda_avance = ws.cell(
                row=config["celda_avance"]["fila"],
                column=config["celda_avance"]["col"],
            )
            if resultado in (None, "", "-") or meta == 0:
                celda_avance.value = None
            else:
                celda_avance.value = float(resultado) / float(meta)
                celda_avance.number_format = "0.00%"

    # Indicador 3: la celda U9:U10 permanece vacía y conserva el plomo de la plantilla.
    ws["U9"] = None

    cantidad_up = _llenar_listado_up(ws_up, df_up, cite)

    # Control de consistencia: detalle UP = valor alcanzado del Indicador 1.
    valor_ind1 = df_res[df_res["Indicador"] == "Indicador 1"].iloc[0]["Valor Alcanzado"]
    if valor_ind1 not in (None, "", "-") and int(float(valor_ind1)) != cantidad_up:
        wb.close()
        raise ValueError(
            f"Inconsistencia para '{cite}': Indicador 1 alcanzado = {valor_ind1}, "
            f"pero el listado UP contiene {cantidad_up} registros."
        )

    output = io.BytesIO()
    wb.save(output)
    wb.close()
    output.seek(0)
    return output.getvalue()


def generar_zip_todas_las_fichas(df, df_metas, df_up, anio, mes_inicio_ui, mes_fin_ui):
    cites = [str(c).strip() for c in df["cite_ut"].dropna().unique()]
    zip_buffer = io.BytesIO()
    mes_ini = mes_inicio_ui[:3].upper()
    mes_fin = mes_fin_ui[:3].upper()
    fichas_generadas = 0

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for cite in cites:
            nombre = (
                f"GDR_{normalizar_nombre_archivo(cite).upper()}_"
                f"{anio}_{mes_ini}_{mes_fin}.xlsx"
            )
            zf.writestr(
                nombre,
                generar_excel_bytes(
                    df, df_metas, df_up, cite, anio, mes_inicio_ui, mes_fin_ui
                ),
            )
            fichas_generadas += 1

    zip_buffer.seek(0)
    return zip_buffer.getvalue(), fichas_generadas
